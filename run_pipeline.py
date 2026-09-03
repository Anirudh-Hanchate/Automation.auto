"""
run_pipeline.py
---------------
Unified CLI runner for the Feedback Automation Pipeline:
1. (Optional) Extract data from scanned PDF forms via Gemini Vision or Claude Vision.
2. Validate and map participant records to sheet columns and criteria rows.
3. Populate the target Excel sheet and save the filled workbook.

Examples:
    # 1. Fill Excel from pre-extracted JSON (fast, offline):
    python run_pipeline.py --json example_extracted_data.json --template blank_template.xlsx --out filled.xlsx

    # 2. Extract from PDF and fill Excel in one step using Gemini:
    python run_pipeline.py --pdf CRE_Feedback.pdf --template blank_template.xlsx --out filled.xlsx --engine gemini

    # 3. Dry-run validation only:
    python run_pipeline.py --json example_extracted_data.json --template blank_template.xlsx --dry-run
"""
import argparse
import json
import os
import sys

import openpyxl
from sheet_map import build_sheet_map
from fill_excel import fill_participant


def run_extraction(pdf_path, json_out, engine="gemini", pages_per_participant=2, model=None, api_key=None, max_participants=None):
    print(f"\n=== Step 1: Extracting feedback from PDF ({engine.upper()}) ===")
    if engine.lower() == "gemini":
        import extract_feedback_gemini as extractor
        default_m = extractor.DEFAULT_MODEL
    elif engine.lower() in ("claude", "anthropic"):
        import extract_feedback as extractor
        default_m = extractor.DEFAULT_MODEL
    else:
        raise ValueError(f"Unknown engine: {engine}. Supported: 'gemini', 'claude'")

    # Set up client and run
    # (Delegating to extractor logic)
    cmd_args = [pdf_path, "--pages-per-participant", str(pages_per_participant), "--out", json_out]
    if model:
        cmd_args.extend(["--model", model])
    if api_key:
        cmd_args.extend(["--api-key", api_key])
    if max_participants:
        cmd_args.extend(["--max-participants", str(max_participants)])

    # Temporarily set sys.argv and call extractor main
    old_argv = sys.argv
    try:
        sys.argv = [f"extract_feedback_{engine}.py"] + cmd_args
        extractor.main()
    finally:
        sys.argv = old_argv


def run_fill(json_path, template_path, sheet_name, out_path, overwrite=False, dry_run=False):
    print(f"\n=== Step 2: Populating Excel Workbook ===")
    print(f"  Source JSON : {json_path}")
    print(f"  Template    : {template_path}")
    print(f"  Target Sheet: {sheet_name}")

    if not os.path.exists(json_path):
        sys.exit(f"Error: JSON file not found: {json_path}")
    if not os.path.exists(template_path):
        sys.exit(f"Error: Template file not found: {template_path}")

    with open(json_path, encoding="utf-8") as f:
        records = json.load(f)

    wb = openpyxl.load_workbook(template_path)
    if sheet_name not in wb.sheetnames:
        sys.exit(f"Error: Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")

    ws = wb[sheet_name]
    smap = build_sheet_map(ws)

    print(f"  Mapped {len(smap.staff_col)} participant columns (from Col {smap.first_data_col} to Col {smap.last_data_col})")
    print(f"  Mapped {len(smap.criteria_rows)} criteria rows across all days")
    if smap.overall_rows:
        print(f"  Mapped {len(smap.overall_rows)} overall feedback rows")
    print(f"  Suggestion 1 Row: {smap.suggestion1_row}, Suggestion 2 Row: {smap.suggestion2_row}\n")

    total_scores = 0
    total_overall = 0
    total_suggestions = 0

    for idx, rec in enumerate(records, 1):
        staff_no = str(rec.get("staff_no", "")).strip()
        name = rec.get("name") or f"Staff #{staff_no}"
        res = fill_participant(ws, smap, rec, overwrite=overwrite)
        total_scores += res["scores"]
        total_overall += res["overall"]
        total_suggestions += res["suggestions"]
        col = smap.staff_col.get(staff_no, "?")
        print(f"  [{idx:02d}/{len(records):02d}] Col {col:>2} | Staff {staff_no:<10} | {name:<25} | {res['scores']:2d} scores, {res['overall']:2d} overall, {res['suggestions']:2d} suggestions")

    print(f"\n=== Summary ===")
    print(f"  Participants processed: {len(records)}")
    print(f"  Score cells written   : {total_scores}")
    print(f"  Overall cells written : {total_overall}")
    print(f"  Suggestions written   : {total_suggestions}")
    print(f"  Total cells populated : {total_scores + total_overall + total_suggestions}")

    if not dry_run:
        target_out = out_path or template_path
        wb.save(target_out)
        print(f"  Workbook successfully saved to: {target_out}")
    else:
        print("  [DRY RUN] Verification successful. Workbook was not modified.")


def main():
    ap = argparse.ArgumentParser(description="Unified Feedback Form Automation Pipeline")
    ap.add_argument("--pdf", default=None, help="Path to input scanned PDF (e.g. CRE_Feedback.pdf)")
    ap.add_argument("--json", default=None, help="Path to extracted data JSON (e.g. example_extracted_data.json or extracted_data.json)")
    ap.add_argument("--template", default="blank_template.xlsx", help="Path to Excel template (default: blank_template.xlsx)")
    ap.add_argument("--sheet", default="Mod-1 Participant's Feedback", help="Target sheet name in template (default: Mod-1 Participant's Feedback)")
    ap.add_argument("--out", default="filled.xlsx", help="Output Excel file path (default: filled.xlsx)")
    ap.add_argument("--engine", choices=["gemini", "claude"], default="gemini", help="Vision AI engine for PDF extraction (default: gemini)")
    ap.add_argument("--model", default=None, help="Specific model ID to use for extraction")
    ap.add_argument("--api-key", default=None, help="API key for extraction")
    ap.add_argument("--pages-per-participant", type=int, default=2, help="Pages per participant in PDF (default: 2)")
    ap.add_argument("--max-participants", type=int, default=None, help="Limit number of participants to extract")
    ap.add_argument("--overwrite", action="store_true", help="Allow overwriting non-empty cells")
    ap.add_argument("--dry-run", action="store_true", help="Validate without writing changes")
    args = ap.parse_args()

    json_path = args.json
    if args.pdf:
        if not json_path:
            json_path = "extracted_data.json"
        if not os.path.exists(json_path) or args.pdf:
            run_extraction(
                pdf_path=args.pdf,
                json_out=json_path,
                engine=args.engine,
                pages_per_participant=args.pages_per_participant,
                model=args.model,
                api_key=args.api_key,
                max_participants=args.max_participants,
            )

    if not json_path:
        if os.path.exists("example_extracted_data.json"):
            json_path = "example_extracted_data.json"
            print(f"No --json or --pdf specified, defaulting to pre-extracted data: {json_path}")
        else:
            sys.exit("Error: Please provide either --pdf <file.pdf> or --json <file.json>.")

    run_fill(
        json_path=json_path,
        template_path=args.template,
        sheet_name=args.sheet,
        out_path=args.out,
        overwrite=args.overwrite,
        dry_run=args.dry_run,
    )


if __name__ == "__main__":
    main()
