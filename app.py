"""
app.py
------
FastAPI backend for Feedback Form to Excel Automation Web Application.
Provides endpoints for file upload, AI vision extraction, Excel filling,
and filled workbook downloads.
"""

import io
import json
import os
import shutil
import tempfile
import time
import uuid
from pathlib import Path
from typing import Optional

import openpyxl
from pypdf import PdfReader
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

import config
from sheet_map import build_sheet_map
from fill_excel import fill_participant
import extract_feedback_gemini as gemini_extractor
import extract_feedback as claude_extractor

BASE_DIR = Path(__file__).parent.resolve()
STATIC_DIR = BASE_DIR / "static"
OUTPUTS_DIR = BASE_DIR / "outputs"
OUTPUTS_DIR.mkdir(exist_ok=True)

DEFAULT_TEMPLATE = BASE_DIR / "blank_template.xlsx"
EXAMPLE_JSON = BASE_DIR / "example_extracted_data.json"

app = FastAPI(title="Feedback Automation API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/api/info")
def get_info():
    sheets = []
    if DEFAULT_TEMPLATE.exists():
        wb = openpyxl.load_workbook(DEFAULT_TEMPLATE, data_only=True)
        sheets = wb.sheetnames

    has_sample_data = EXAMPLE_JSON.exists()
    gemini_env = bool(config.GEMINI_API_KEY or os.environ.get("GEMINI_API_KEY"))
    anthropic_env = bool(config.ANTHROPIC_API_KEY or os.environ.get("ANTHROPIC_API_KEY"))

    return {
        "status": "ready",
        "default_template": DEFAULT_TEMPLATE.name,
        "available_sheets": sheets,
        "has_sample_data": has_sample_data,
        "has_gemini_key": gemini_env,
        "has_anthropic_key": anthropic_env,
    }

@app.post("/api/process-stream")
async def process_feedback_stream(
    pdf_file: Optional[UploadFile] = File(None),
    template_file: Optional[UploadFile] = File(None),
    mode: str = Form("demo"),
    api_key: Optional[str] = Form(None),
    model: Optional[str] = Form(None),
    sheet_name: str = Form(config.DEFAULT_SHEET_NAME),
    pages_per_participant: int = Form(config.DEFAULT_PAGES_PER_PARTICIPANT),
    overwrite: bool = Form(True),
):
    job_id = str(uuid.uuid4())[:8]
    temp_dir = Path(tempfile.mkdtemp(prefix=f"feedback_{job_id}_"))

    # Read uploaded file contents before entering generator
    pdf_bytes = None
    pdf_filename = None
    if pdf_file and pdf_file.filename:
        pdf_bytes = await pdf_file.read()
        pdf_filename = pdf_file.filename

    template_bytes = None
    if template_file and template_file.filename:
        template_bytes = await template_file.read()

    def event_stream():
        try:
            yield f"data: {json.dumps({'type': 'log', 'message': '> Initializing automation pipeline...'})}\n\n"

            # 1. Template preparation
            template_path = temp_dir / "template.xlsx"
            if template_bytes:
                with open(template_path, "wb") as f:
                    f.write(template_bytes)
                yield f"data: {json.dumps({'type': 'log', 'message': '> Loaded custom Excel template.'})}\n\n"
            elif DEFAULT_TEMPLATE.exists():
                shutil.copy(DEFAULT_TEMPLATE, template_path)
                yield f"data: {json.dumps({'type': 'log', 'message': '> Using default template: blank_template.xlsx'})}\n\n"
            else:
                yield f"data: {json.dumps({'type': 'error', 'message': 'No template file found.'})}\n\n"
                return

            records = []
            if mode == "demo":
                yield f"data: {json.dumps({'type': 'log', 'message': '> Loading pre-extracted demo records (19 participants)...'})}\n\n"
                if not EXAMPLE_JSON.exists():
                    yield f"data: {json.dumps({'type': 'error', 'message': 'Example data file not found on server.'})}\n\n"
                    return
                with open(EXAMPLE_JSON, encoding="utf-8") as f:
                    records = json.load(f)
                yield f"data: {json.dumps({'type': 'progress', 'current': len(records), 'total': len(records), 'percent': 50})}\n\n"

            else:
                if not pdf_bytes:
                    yield f"data: {json.dumps({'type': 'error', 'message': 'Please upload a PDF feedback form.'})}\n\n"
                    return

                pdf_path = temp_dir / (pdf_filename or "upload.pdf")
                with open(pdf_path, "wb") as f:
                    f.write(pdf_bytes)

                reader = PdfReader(str(pdf_path))
                n_pages = len(reader.pages)
                step = pages_per_participant
                chunk_starts = list(range(0, n_pages, step))
                total_chunks = len(chunk_starts)

                yield f"data: {json.dumps({'type': 'log', 'message': f'> Uploaded PDF: {pdf_filename} ({n_pages} pages, {total_chunks} participants found)'})}\n\n"

                if mode == "gemini":
                    key = (api_key or "").strip() or config.get_gemini_key()
                    if not key:
                        yield f"data: {json.dumps({'type': 'error', 'message': 'Gemini API key is missing. Add it to .env or paste it in Settings.'})}\n\n"
                        return

                    client = gemini_extractor.genai.Client(api_key=key, vertexai=False)
                    m_name = model or config.DEFAULT_GEMINI_MODEL
                    workers_count = 5
                    yield f"data: {json.dumps({'type': 'log', 'message': f'> ⚡ Launching Parallel Vision Engine ({workers_count} concurrent workers, model: {m_name})...'})}\n\n"

                    import queue
                    from concurrent.futures import ThreadPoolExecutor, as_completed

                    event_q = queue.Queue()
                    extracted_records = [None] * total_chunks

                    # Pre-encode all chunks in memory for fast parallel dispatch
                    encoded_chunks = []
                    for idx, start in enumerate(chunk_starts):
                        idxs = list(range(start, min(start + step, n_pages)))
                        label = f"Participant {idx+1}/{total_chunks} (pages {idxs[0]+1}-{idxs[-1]+1})"
                        chunk_pdf = gemini_extractor.encode_pdf_chunk(reader, idxs)
                        encoded_chunks.append((idx, chunk_pdf, label))

                    def _gemini_worker(item):
                        c_idx, pdf_data, lbl = item
                        try:
                            rec = gemini_extractor.call_gemini(client, pdf_data, lbl, m_name)
                            extracted_records[c_idx] = rec
                            event_q.put(("success", c_idx, rec, lbl))
                        except Exception as exc:
                            event_q.put(("error", c_idx, str(exc), lbl))

                    # Start thread pool executor
                    completed_count = 0
                    with ThreadPoolExecutor(max_workers=workers_count) as executor:
                        for chunk_item in encoded_chunks:
                            executor.submit(_gemini_worker, chunk_item)

                        while completed_count < total_chunks:
                            try:
                                status, c_idx, payload, lbl = event_q.get(timeout=30.0)
                                completed_count += 1
                                pct = int((completed_count / total_chunks) * 75)

                                if status == "success":
                                    s_no = payload.get("staff_no", "Unknown")
                                    p_name = payload.get("name", f"Staff #{s_no}")
                                    yield f"data: {json.dumps({'type': 'progress', 'current': completed_count, 'total': total_chunks, 'percent': pct})}\n\n"
                                    yield f"data: {json.dumps({'type': 'log', 'message': f'> [✓] [{completed_count:02d}/{total_chunks:02d}] Extracted: {p_name} (Staff No: {s_no})'})}\n\n"
                                else:
                                    yield f"data: {json.dumps({'type': 'log', 'message': f'> [WARNING] {lbl} warning: {payload}'})}\n\n"
                            except queue.Empty:
                                break

                    records = [r for r in extracted_records if r is not None]

                elif mode == "claude":
                    key = (api_key or "").strip() or config.get_anthropic_key()
                    if not key:
                        yield f"data: {json.dumps({'type': 'error', 'message': 'Anthropic API key is missing. Add it to .env or paste it in Settings.'})}\n\n"
                        return

                    client = claude_extractor.anthropic.Anthropic(api_key=key)
                    m_name = model or config.DEFAULT_CLAUDE_MODEL
                    yield f"data: {json.dumps({'type': 'log', 'message': f'> ⚡ Starting AI Vision extraction using Claude {m_name}...'})}\n\n"

                    import queue
                    from concurrent.futures import ThreadPoolExecutor

                    event_q = queue.Queue()
                    extracted_records = [None] * total_chunks
                    encoded_chunks = []
                    for idx, start in enumerate(chunk_starts):
                        idxs = list(range(start, min(start + step, n_pages)))
                        label = f"Participant {idx+1}/{total_chunks} (pages {idxs[0]+1}-{idxs[-1]+1})"
                        chunk_pdf = claude_extractor.encode_pdf_chunk(reader, idxs)
                        encoded_chunks.append((idx, chunk_pdf, label))

                    def _claude_worker(item):
                        c_idx, pdf_data, lbl = item
                        try:
                            rec = claude_extractor.call_claude(client, pdf_data, lbl, m_name)
                            extracted_records[c_idx] = rec
                            event_q.put(("success", c_idx, rec, lbl))
                        except Exception as exc:
                            event_q.put(("error", c_idx, str(exc), lbl))

                    completed_count = 0
                    with ThreadPoolExecutor(max_workers=4) as executor:
                        for chunk_item in encoded_chunks:
                            executor.submit(_claude_worker, chunk_item)

                        while completed_count < total_chunks:
                            try:
                                status, c_idx, payload, lbl = event_q.get(timeout=30.0)
                                completed_count += 1
                                pct = int((completed_count / total_chunks) * 75)

                                if status == "success":
                                    s_no = payload.get("staff_no", "Unknown")
                                    p_name = payload.get("name", f"Staff #{s_no}")
                                    yield f"data: {json.dumps({'type': 'progress', 'current': completed_count, 'total': total_chunks, 'percent': pct})}\n\n"
                                    yield f"data: {json.dumps({'type': 'log', 'message': f'> [✓] [{completed_count:02d}/{total_chunks:02d}] Extracted: {p_name} (Staff No: {s_no})'})}\n\n"
                                else:
                                    yield f"data: {json.dumps({'type': 'log', 'message': f'> [WARNING] {lbl} warning: {payload}'})}\n\n"
                            except queue.Empty:
                                break

                    records = [r for r in extracted_records if r is not None]

                else:
                    yield f"data: {json.dumps({'type': 'error', 'message': f'Unsupported mode: {mode}'})}\n\n"
                    return

            if not records:
                yield f"data: {json.dumps({'type': 'error', 'message': 'No participant records could be extracted.'})}\n\n"
                return

            # Save a backup of extracted records so data is never lost
            try:
                backup_json = OUTPUTS_DIR / "last_extracted_data.json"
                with open(backup_json, "w", encoding="utf-8") as bf:
                    json.dump(records, bf, indent=2, ensure_ascii=False)
            except Exception:
                pass

            # Fill Excel
            msg = f"> Populating Excel sheet '{sheet_name}'..."
            yield f"data: {json.dumps({'type': 'log', 'message': msg})}\n\n"
            wb = openpyxl.load_workbook(template_path)
            if sheet_name not in wb.sheetnames:
                err_msg = f"Sheet '{sheet_name}' not found in template. Available: {wb.sheetnames}"
                yield f"data: {json.dumps({'type': 'error', 'message': err_msg})}\n\n"
                return

            ws = wb[sheet_name]
            smap = build_sheet_map(ws)

            total_scores = 0
            total_overall = 0
            total_suggestions = 0
            participant_summaries = []

            for rec in records:
                staff_no = str(rec.get("staff_no", "")).strip()
                name = rec.get("name") or f"Staff #{staff_no}"
                res = fill_participant(ws, smap, rec, overwrite=overwrite, auto_add_column=True)
                total_scores += res["scores"]
                total_overall += res["overall"]
                total_suggestions += res["suggestions"]
                participant_summaries.append({
                    "staff_no": staff_no,
                    "name": name,
                    "scores_filled": res["scores"],
                    "overall_filled": res["overall"],
                    "suggestions_filled": res["suggestions"],
                    "suggestion_1": rec.get("suggestion_1", ""),
                    "suggestion_2": rec.get("suggestion_2", ""),
                })

            out_filename = f"filled_feedback_{job_id}.xlsx"
            out_path = OUTPUTS_DIR / out_filename
            wb.save(out_path)

            yield f"data: {json.dumps({'type': 'progress', 'current': len(records), 'total': len(records), 'percent': 100})}\n\n"
            yield f"data: {json.dumps({'type': 'complete', 'job_id': job_id, 'filename': out_filename, 'download_url': f'/api/download/{out_filename}', 'summary': {'participants_count': len(records), 'total_scores': total_scores, 'total_overall': total_overall, 'total_suggestions': total_suggestions, 'total_cells': total_scores + total_overall + total_suggestions}, 'participants': participant_summaries})}\n\n"

        except Exception as exc:
            yield f"data: {json.dumps({'type': 'error', 'message': str(exc)})}\n\n"
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    return StreamingResponse(event_stream(), media_type="text/event-stream")


@app.post("/api/process")
async def process_feedback(
    pdf_file: Optional[UploadFile] = File(None),
    template_file: Optional[UploadFile] = File(None),
    mode: str = Form("demo"),  # "demo", "gemini", "claude"
    api_key: Optional[str] = Form(None),
    model: Optional[str] = Form(None),
    sheet_name: str = Form(config.DEFAULT_SHEET_NAME),
    pages_per_participant: int = Form(config.DEFAULT_PAGES_PER_PARTICIPANT),
    overwrite: bool = Form(True),
):
    job_id = str(uuid.uuid4())[:8]
    temp_dir = Path(tempfile.mkdtemp(prefix=f"feedback_{job_id}_"))

    try:
        # 1. Determine Template
        template_path = temp_dir / "template.xlsx"
        if template_file and template_file.filename:
            with open(template_path, "wb") as f:
                shutil.copyfileobj(template_file.file, f)
        elif DEFAULT_TEMPLATE.exists():
            shutil.copy(DEFAULT_TEMPLATE, template_path)
        else:
            raise HTTPException(status_code=400, detail="No template file available.")

        # 2. Extract or Load Records
        records = []
        if mode == "demo":
            if not EXAMPLE_JSON.exists():
                raise HTTPException(status_code=404, detail="Example extracted data not found on server.")
            with open(EXAMPLE_JSON, encoding="utf-8") as f:
                records = json.load(f)
        else:
            if not pdf_file or not pdf_file.filename:
                raise HTTPException(status_code=400, detail="Please upload a PDF feedback file to extract.")

            pdf_path = temp_dir / pdf_file.filename
            with open(pdf_path, "wb") as f:
                shutil.copyfileobj(pdf_file.file, f)

            reader = PdfReader(str(pdf_path))
            n_pages = len(reader.pages)
            step = pages_per_participant

            if mode == "gemini":
                key = (api_key or "").strip() or config.get_gemini_key()
                if not key:
                    raise HTTPException(
                        status_code=400,
                        detail="Gemini API key is not configured. Please paste your Google AI Studio API key (starts with 'AIzaSy...') in config.py or in the settings panel."
                    )
                if key.startswith("AQ."):
                    raise HTTPException(
                        status_code=400,
                        detail="The provided key is an IDE/OAuth session token (starts with 'AQ.'), not a Google AI Studio API key. Please get a free API key (starts with 'AIzaSy...') from https://aistudio.google.com/apikey"
                    )
                client = gemini_extractor.genai.Client(api_key=key, vertexai=False)
                m_name = model or config.DEFAULT_GEMINI_MODEL

                for idx, start in enumerate(range(0, n_pages, step), 1):
                    idxs = list(range(start, min(start + step, n_pages)))
                    label = f"Participant {idx} (pages {idxs[0]+1}-{idxs[-1]+1})"
                    chunk_pdf = gemini_extractor.encode_pdf_chunk(reader, idxs)
                    try:
                        rec = gemini_extractor.call_gemini(client, chunk_pdf, label, m_name)
                        records.append(rec)
                    except Exception as err:
                        print(f"Warning: could not extract {label}: {err}")
                    time.sleep(2.5)

            elif mode == "claude":
                key = (api_key or "").strip() or config.get_anthropic_key()
                if not key:
                    raise HTTPException(
                        status_code=400,
                        detail="Anthropic API key is not configured. Please paste your Anthropic API key (starts with 'sk-ant-...') in config.py or in the settings panel."
                    )
                client = claude_extractor.anthropic.Anthropic(api_key=key)

                m_name = model or config.DEFAULT_CLAUDE_MODEL

                for idx, start in enumerate(range(0, n_pages, step), 1):
                    idxs = list(range(start, min(start + step, n_pages)))
                    label = f"Participant {idx} (pages {idxs[0]+1}-{idxs[-1]+1})"
                    chunk_pdf = claude_extractor.encode_pdf_chunk(reader, idxs)
                    try:
                        rec = claude_extractor.call_claude(client, chunk_pdf, label, m_name)
                        records.append(rec)
                    except Exception as err:
                        print(f"Warning: could not extract {label}: {err}")

            else:
                raise HTTPException(status_code=400, detail=f"Unsupported mode: {mode}")

        # 3. Fill Excel Workbook
        wb = openpyxl.load_workbook(template_path)
        if sheet_name not in wb.sheetnames:
            raise HTTPException(status_code=400, detail=f"Sheet '{sheet_name}' not found in template. Available: {wb.sheetnames}")

        ws = wb[sheet_name]
        smap = build_sheet_map(ws)

        total_scores = 0
        total_overall = 0
        total_suggestions = 0
        participant_summaries = []

        for rec in records:
            staff_no = str(rec.get("staff_no", "")).strip()
            name = rec.get("name") or f"Staff #{staff_no}"
            res = fill_participant(ws, smap, rec, overwrite=overwrite, auto_add_column=True)
            total_scores += res["scores"]
            total_overall += res["overall"]
            total_suggestions += res["suggestions"]
            participant_summaries.append({
                "staff_no": staff_no,
                "name": name,
                "scores_filled": res["scores"],
                "overall_filled": res["overall"],
                "suggestions_filled": res["suggestions"],
                "suggestion_1": rec.get("suggestion_1", ""),
                "suggestion_2": rec.get("suggestion_2", ""),
            })

        # 4. Save Output
        out_filename = f"filled_feedback_{job_id}.xlsx"
        out_path = OUTPUTS_DIR / out_filename
        wb.save(out_path)

        return {
            "status": "success",
            "job_id": job_id,
            "filename": out_filename,
            "download_url": f"/api/download/{out_filename}",
            "summary": {
                "participants_count": len(records),
                "total_scores": total_scores,
                "total_overall": total_overall,
                "total_suggestions": total_suggestions,
                "total_cells": total_scores + total_overall + total_suggestions,
            },
            "participants": participant_summaries,
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

@app.get("/api/download/{filename}")
def download_file(filename: str):
    file_path = OUTPUTS_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found or expired.")
    return FileResponse(
        path=file_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

# Serve Static Frontend Files
app.mount("/", StaticFiles(directory=str(STATIC_DIR), html=True), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
